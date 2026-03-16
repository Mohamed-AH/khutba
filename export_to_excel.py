#!/usr/bin/env python3
"""
export_to_excel.py — Export khutba archive to Excel table.

Columns:  S.No | TelegramFileName | Type | SeriesName | Sheikh | DateInGreg | Category

Usage:
    python export_to_excel.py

Expects khutba_archive.json (and optionally filename.txt) in the same directory.
Outputs khutba_archive.xlsx.
"""

import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ARCHIVE_JSON = Path("khutba_archive.json")
FILENAME_TXT = Path("filename.txt")
OUTPUT_XLSX  = Path("khutba_archive.xlsx")

SHEIKH = "حسن بن محمد منصور الدغريري"


def load_filenames() -> dict[int, str]:
    """Parse filename.txt → {index: basename}  e.g. {9: '009_من_أحكام_النكاح_1.m4a'}
    Handles Windows-style backslash paths when run on Linux."""
    names: dict[int, str] = {}
    if not FILENAME_TXT.exists():
        return names
    for line in FILENAME_TXT.read_text(encoding="utf-8").splitlines():
        line = line.strip().strip('"')
        if not line:
            continue
        # Extract basename from either / or \ separated paths
        basename = re.split(r'[/\\]', line)[-1]
        m = re.match(r'^(\d+)_', basename)
        if m:
            names[int(m.group(1))] = basename
    return names


def build_rows() -> list[dict]:
    with open(ARCHIVE_JSON, encoding="utf-8") as f:
        records = json.load(f)

    filenames = load_filenames()
    rows = []

    for sno, rec in enumerate(records, start=1):
        idx      = rec["index"]
        date_str = rec.get("date") or ""
        title    = rec.get("title") or ""

        tg_name = filenames.get(idx, "")

        # Date formatted as DD.MM.YYYY
        try:
            dt = datetime.fromisoformat(date_str)
            date_greg = dt.strftime("%d.%m.%Y")
        except Exception:
            date_greg = ""

        rows.append({
            "S.No":             sno,
            "TelegramFileName": tg_name,
            "Type":             "Khutba",
            "SeriesName":       title,
            "Sheikh":           SHEIKH,
            "DateInGreg":       date_greg,
            "Category":         "Khutba",
        })

    return rows


# ── Excel styling ────────────────────────────────────────────────────────────

HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")   # dark blue
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
ROW_FONT     = Font(size=11)
ROW_FONT_RTL = Font(size=11)

COLUMNS = ["S.No", "TelegramFileName", "Type", "SeriesName", "Sheikh", "DateInGreg", "Category"]
COL_WIDTHS = {
    "S.No":             6,
    "TelegramFileName": 34,
    "Type":             16,
    "SeriesName":       52,
    "Sheikh":           32,
    "DateInGreg":       14,
    "Category":         14,
}


def write_excel(rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Khutba Archive"
    ws.sheet_view.rightToLeft = False

    # Header row
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=False)

    ws.row_dimensions[1].height = 20

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            value = row[col_name]
            cell  = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = ROW_FONT

            # RTL alignment for Arabic columns
            if col_name in ("SeriesName", "Sheikh"):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_name == "S.No":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        # Zebra striping
        if row_idx % 2 == 0:
            fill = PatternFill("solid", fgColor="DCE6F1")
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=c).fill = fill

    # Column widths
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS[col_name]

    # Freeze header
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_XLSX)
    print(f"Saved → {OUTPUT_XLSX}  ({len(rows)} rows)")


def main():
    if not ARCHIVE_JSON.exists():
        print(f"ERROR: {ARCHIVE_JSON} not found.")
        return
    rows = build_rows()
    write_excel(rows)

    # Quick console preview (first 3 rows)
    print("\nPreview (first 3 rows):")
    print(f"{'─'*90}")
    for r in rows[:3]:
        print(f"  {r['S.No']:>3}.  {r['TelegramFileName']:<36} {r['Type']:<16} {r['DateInGreg']}")
        print(f"       SeriesName : {r['SeriesName']}")
        print(f"       Sheikh     : {r['Sheikh']}\n")


if __name__ == "__main__":
    main()
